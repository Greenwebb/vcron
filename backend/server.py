from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response
from fastapi.responses import JSONResponse, StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
import asyncio
from pathlib import Path
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional
import uuid
from datetime import datetime, timezone, timedelta
import bcrypt
import jwt
import httpx
import resend
from io import BytesIO
import openpyxl

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# JWT Configuration
JWT_SECRET = os.environ.get('JWT_SECRET', 'vchron-super-secret-key-2024')
JWT_ALGORITHM = "HS256"
JWT_EXPIRATION_DAYS = 7

# Resend Configuration
resend.api_key = os.environ.get('RESEND_API_KEY', '')
SENDER_EMAIL = os.environ.get('SENDER_EMAIL', 'onboarding@resend.dev')
BACKUP_EMAIL = os.environ.get('BACKUP_EMAIL', 'northmkushidistrict@gmail.com')

# Create the main app
app = FastAPI(title="V-Chron API", description="Healthcare Attendance Tracking System")

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# ===================== PROVINCES LIST =====================
PROVINCES = [
    {"id": 1, "name": "Central Province"},
    {"id": 2, "name": "Copperbelt Province"},
    {"id": 3, "name": "Eastern Province"},
    {"id": 4, "name": "Luapula Province"},
    {"id": 5, "name": "Lusaka Province"},
    {"id": 6, "name": "Muchinga Province"},
    {"id": 7, "name": "Northern Province"},
    {"id": 8, "name": "North-Western Province"},
    {"id": 9, "name": "Southern Province"},
    {"id": 10, "name": "Western Province"}
]

# ===================== DISTRICTS BY PROVINCE =====================
DISTRICTS = {
    "Central Province": [
        "Chibombo",
        "Chisamba",
        "Chitambo",
        "Kabwe",
        "Kapiri Mposhi",
        "Luano",
        "Mkushi",
        "Mumbwa",
        "Ngabwe",
        "Serenje",
        "Shibuyunji"
    ]
}

# ===================== FACILITIES BY DISTRICT =====================
FACILITIES_BY_DISTRICT = {
    "Mkushi": [
        "Chalata Rural Health Centre",
        "Changilo Health Post",
        "Chibwemukunga Health Post",
        "Nambo Rural Health Post",
        "Ntekete Health Post",
        "Chibefwe Health Centre",
        "Milombwe Health Post",
        "Tazara Health Post",
        "Mulundu Health Post",
        "Chitina Health Post",
        "Chikabile Health Post",
        "Kabengeshi Rural Health Post",
        "Kakushi Health Post",
        "Luanshimba Rural Health Centre (Mkushi)",
        "Mboboli Rural Health Post",
        "Mulungwe Rural Health Centre",
        "Matuku Health Post",
        "Miloso Health Post",
        "Momboshi Health Post",
        "Twatasha Health Post",
        "Fibanga Health Post",
        "Munsakamba Health Post",
        "Katuba Health Post",
        "Nkulumashiba Rural Health Post",
        "Chisanga Rural Health Centre",
        "Musofu Rural Health Centre",
        "Upper Musofu Health Post",
        "Chengelo Clinic",
        "Kasalamakanga Health Post",
        "Nkolonga Farm Clinic",
        "Nkolonga Health Post",
        "Chine Rural Health Post",
        "Fiwila Rural Health Centre",
        "Kalubula Rural Health Post",
        "Nyenje Health Post",
        "Shaibila Health Post",
        "Kasokota Health Post",
        "Mikunku Rural Health Centre",
        "Nkumbi College Clinic",
        "Nkumbi Rural Health Centre",
        "Lilanda Health Post",
        "Malubila Health Post",
        "Mankanda Health Post",
        "Milele Health Post",
        "Nshinso Rural Health Centre",
        "Upper Lusemfwa Health Post",
        "Mkushi District Hospital"
    ]
}

# Legacy FACILITIES list for backward compatibility
FACILITIES = FACILITIES_BY_DISTRICT.get("Mkushi", [])

POSITIONS = [
    "Nurse",
    "Clinical Officer",
    "Environmental Health Technician",
    "Medical Doctor",
    "Pharmacist",
    "Laboratory Technician",
    "Radiographer",
    "Physiotherapist",
    "Midwife",
    "Community Health Worker",
    "Administrative Staff",
    "Other"
]

AREAS_OF_ALLOCATION = [
    "Facility",
    "Outreach"
]

# ===================== DEFAULT SHIFT TIMES =====================
DEFAULT_SHIFTS = {
    "morning": {"start": "06:00", "end": "14:00"},
    "afternoon": {"start": "14:00", "end": "22:00"},
    "night": {"start": "22:00", "end": "06:00"},
    "four_off": {"start": "07:00", "end": "19:00"}
}

# ===================== PYDANTIC MODELS =====================

class UserBase(BaseModel):
    email: EmailStr
    name: str
    phone_number: Optional[str] = None
    position: Optional[str] = None
    province: Optional[str] = None
    district: Optional[str] = None
    facility: Optional[str] = None
    area_of_allocation: Optional[str] = None
    picture: Optional[str] = None
    role: str = "user"  # user, admin, superuser
    assigned_scope: Optional[dict] = None  # For admins: {"type": "facility/district/province", "value": "name"}

class UserCreate(BaseModel):
    email: EmailStr
    password: str
    name: str
    phone_number: str
    position: str
    province: str
    district: str
    facility: str

class UserLogin(BaseModel):
    email: EmailStr
    password: str

class UserResponse(BaseModel):
    user_id: str
    email: str
    name: str
    phone_number: Optional[str] = None
    position: Optional[str] = None
    province: Optional[str] = None
    district: Optional[str] = None
    facility: Optional[str] = None
    area_of_allocation: Optional[str] = None
    picture: Optional[str] = None
    role: str = "user"
    assigned_scope: Optional[dict] = None
    assigned_shift: Optional[str] = None
    created_at: Optional[str] = None

class UserUpdate(BaseModel):
    name: Optional[str] = None
    phone_number: Optional[str] = None
    position: Optional[str] = None
    province: Optional[str] = None
    district: Optional[str] = None
    facility: Optional[str] = None
    area_of_allocation: Optional[str] = None
    role: Optional[str] = None
    assigned_scope: Optional[dict] = None
    assigned_shift: Optional[str] = None

class FacilityCreate(BaseModel):
    name: str
    district: str
    province: str

class FacilityUpdate(BaseModel):
    name: Optional[str] = None
    district: Optional[str] = None
    province: Optional[str] = None

class ShiftConfig(BaseModel):
    morning_start: str = "06:00"
    morning_end: str = "14:00"
    afternoon_start: str = "14:00"
    afternoon_end: str = "22:00"
    night_start: str = "22:00"
    night_end: str = "06:00"
    four_off_start: str = "07:00"
    four_off_end: str = "19:00"
    grace_period_minutes: int = 15

class PasswordReset(BaseModel):
    new_password: str

class AttendanceRecord(BaseModel):
    attendance_id: str
    user_id: str
    user_name: str
    position: str
    facility: str
    area_of_allocation: Optional[str] = None
    action: str  # "login" or "logout"
    timestamp: str
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    synced: bool = True

class AttendanceCreate(BaseModel):
    action: str  # "login" or "logout"
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    area_of_allocation: Optional[str] = None  # Facility or Outreach
    offline_id: Optional[str] = None  # For offline sync

class OfflineAttendanceSync(BaseModel):
    records: List[AttendanceCreate]

class TokenResponse(BaseModel):
    access_token: str
    token_type: str = "bearer"
    user: UserResponse

# ===================== HELPER FUNCTIONS =====================

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')

def verify_password(password: str, hashed: str) -> bool:
    return bcrypt.checkpw(password.encode('utf-8'), hashed.encode('utf-8'))

def create_jwt_token(user_id: str) -> str:
    expiration = datetime.now(timezone.utc) + timedelta(days=JWT_EXPIRATION_DAYS)
    payload = {
        "user_id": user_id,
        "exp": expiration,
        "iat": datetime.now(timezone.utc)
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)

def decode_jwt_token(token: str) -> dict:
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        return payload
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token has expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")

async def get_current_user(request: Request) -> dict:
    # Check cookie first
    session_token = request.cookies.get("session_token")
    
    # Then check Authorization header
    if not session_token:
        auth_header = request.headers.get("Authorization")
        if auth_header and auth_header.startswith("Bearer "):
            session_token = auth_header.split(" ")[1]
    
    if not session_token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    
    # Check if it's a JWT token (for email/password auth)
    try:
        payload = decode_jwt_token(session_token)
        user_id = payload.get("user_id")
        user = await db.users.find_one({"user_id": user_id}, {"_id": 0})
        if not user:
            raise HTTPException(status_code=401, detail="User not found")
        return user
    except Exception:
        pass
    
    # Check if it's a session token (for Google OAuth)
    session = await db.user_sessions.find_one({"session_token": session_token}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=401, detail="Invalid session")
    
    # Check expiration
    expires_at = session.get("expires_at")
    if isinstance(expires_at, str):
        expires_at = datetime.fromisoformat(expires_at)
    if expires_at.tzinfo is None:
        expires_at = expires_at.replace(tzinfo=timezone.utc)
    if expires_at < datetime.now(timezone.utc):
        raise HTTPException(status_code=401, detail="Session expired")
    
    user = await db.users.find_one({"user_id": session["user_id"]}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=401, detail="User not found")
    
    return user

async def get_admin_user(request: Request) -> dict:
    user = await get_current_user(request)
    if user.get("role") not in ["admin", "superuser"]:
        raise HTTPException(status_code=403, detail="Admin access required")
    return user

async def get_superuser(request: Request) -> dict:
    user = await get_current_user(request)
    if user.get("role") != "superuser":
        raise HTTPException(status_code=403, detail="Super user access required")
    return user

# ===================== AUTH ROUTES =====================

@api_router.post("/auth/register", response_model=TokenResponse)
async def register(user_data: UserCreate):
    # Check if user exists
    existing = await db.users.find_one({"email": user_data.email}, {"_id": 0})
    if existing:
        raise HTTPException(status_code=400, detail="Email already registered")
    
    # Validate province
    province_names = [p["name"] for p in PROVINCES]
    if user_data.province not in province_names:
        raise HTTPException(status_code=400, detail="Invalid province")
    
    # Validate district (check hardcoded + DB)
    hardcoded_districts = DISTRICTS.get(user_data.province, [])
    db_facs = await db.facilities.find({"province": user_data.province}, {"_id": 0, "district": 1}).to_list(1000)
    db_districts = [f["district"] for f in db_facs if f.get("district")]
    all_districts = set(hardcoded_districts + db_districts)
    if user_data.district not in all_districts:
        raise HTTPException(status_code=400, detail="Invalid district for selected province")
    
    # Validate facility (check hardcoded + DB)
    hardcoded_facilities = FACILITIES_BY_DISTRICT.get(user_data.district, [])
    db_facs2 = await db.facilities.find({"district": user_data.district}, {"_id": 0, "name": 1}).to_list(1000)
    db_facility_names = [f["name"] for f in db_facs2 if f.get("name")]
    all_facilities = set(hardcoded_facilities + db_facility_names)
    if user_data.facility not in all_facilities:
        raise HTTPException(status_code=400, detail="Invalid facility for selected district")
    
    # Validate position
    if user_data.position not in POSITIONS:
        raise HTTPException(status_code=400, detail="Invalid position")
    
    user_id = f"user_{uuid.uuid4().hex[:12]}"
    hashed_password = hash_password(user_data.password)
    
    user_doc = {
        "user_id": user_id,
        "email": user_data.email,
        "name": user_data.name,
        "phone_number": user_data.phone_number,
        "password": hashed_password,
        "position": user_data.position,
        "province": user_data.province,
        "district": user_data.district,
        "facility": user_data.facility,
        "picture": None,
        "role": "user",
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    
    await db.users.insert_one(user_doc)
    
    # Generate token
    token = create_jwt_token(user_id)
    
    return TokenResponse(
        access_token=token,
        user=UserResponse(
            user_id=user_id,
            email=user_data.email,
            name=user_data.name,
            phone_number=user_data.phone_number,
            position=user_data.position,
            province=user_data.province,
            district=user_data.district,
            facility=user_data.facility,
            role="user",
            created_at=user_doc["created_at"]
        )
    )

@api_router.post("/auth/login", response_model=TokenResponse)
async def login(credentials: UserLogin, response: Response):
    user = await db.users.find_one({"email": credentials.email}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=401, detail="Invalid email or password")
    
    if not user.get("password"):
        raise HTTPException(status_code=401, detail="Please use Google sign-in for this account")
    
    if not verify_password(credentials.password, user["password"]):
        raise HTTPException(status_code=401, detail="Invalid email or password")
    
    token = create_jwt_token(user["user_id"])
    
    # Set cookie
    response.set_cookie(
        key="session_token",
        value=token,
        httponly=True,
        secure=True,
        samesite="none",
        path="/",
        max_age=JWT_EXPIRATION_DAYS * 24 * 60 * 60
    )
    
    return TokenResponse(
        access_token=token,
        user=UserResponse(
            user_id=user["user_id"],
            email=user["email"],
            name=user["name"],
            phone_number=user.get("phone_number"),
            position=user.get("position"),
            province=user.get("province"),
            district=user.get("district"),
            facility=user.get("facility"),
            area_of_allocation=user.get("area_of_allocation"),
            picture=user.get("picture"),
            role=user.get("role", "user"),
            created_at=user.get("created_at")
        )
    )

@api_router.post("/auth/session")
async def process_session(request: Request, response: Response):
    """Process Google OAuth session_id and create session"""
    body = await request.json()
    session_id = body.get("session_id")
    
    if not session_id:
        raise HTTPException(status_code=400, detail="session_id required")
    
    # Call Emergent Auth to get session data
    async with httpx.AsyncClient() as client:
        try:
            auth_response = await client.get(
                "https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data",
                headers={"X-Session-ID": session_id},
                timeout=10.0
            )
            if auth_response.status_code != 200:
                raise HTTPException(status_code=401, detail="Invalid session")
            
            session_data = auth_response.json()
        except httpx.RequestError as e:
            logger.error(f"Error calling auth service: {e}")
            raise HTTPException(status_code=500, detail="Authentication service error")
    
    email = session_data.get("email")
    name = session_data.get("name")
    picture = session_data.get("picture")
    session_token = session_data.get("session_token")
    
    # Check if user exists
    user = await db.users.find_one({"email": email}, {"_id": 0})
    
    if user:
        # Update existing user
        await db.users.update_one(
            {"email": email},
            {"$set": {"name": name, "picture": picture}}
        )
        user_id = user["user_id"]
    else:
        # Create new user (needs to complete registration)
        user_id = f"user_{uuid.uuid4().hex[:12]}"
        user = {
            "user_id": user_id,
            "email": email,
            "name": name,
            "picture": picture,
            "position": None,
            "facility": None,
            "role": "user",
            "created_at": datetime.now(timezone.utc).isoformat()
        }
        await db.users.insert_one(user)
    
    # Store session
    expires_at = datetime.now(timezone.utc) + timedelta(days=7)
    await db.user_sessions.insert_one({
        "user_id": user_id,
        "session_token": session_token,
        "expires_at": expires_at.isoformat(),
        "created_at": datetime.now(timezone.utc).isoformat()
    })
    
    # Set cookie
    response.set_cookie(
        key="session_token",
        value=session_token,
        httponly=True,
        secure=True,
        samesite="none",
        path="/",
        max_age=7 * 24 * 60 * 60
    )
    
    # Get updated user
    user = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    
    return {
        "user_id": user["user_id"],
        "email": user["email"],
        "name": user["name"],
        "phone_number": user.get("phone_number"),
        "position": user.get("position"),
        "province": user.get("province"),
        "district": user.get("district"),
        "facility": user.get("facility"),
        "area_of_allocation": user.get("area_of_allocation"),
        "picture": user.get("picture"),
        "role": user.get("role", "user"),
        "needs_registration": user.get("position") is None or user.get("facility") is None
    }

@api_router.get("/auth/me", response_model=UserResponse)
async def get_me(user: dict = Depends(get_current_user)):
    return UserResponse(
        user_id=user["user_id"],
        email=user["email"],
        name=user["name"],
        phone_number=user.get("phone_number"),
        position=user.get("position"),
        province=user.get("province"),
        district=user.get("district"),
        facility=user.get("facility"),
        area_of_allocation=user.get("area_of_allocation"),
        picture=user.get("picture"),
        role=user.get("role", "user"),
        created_at=user.get("created_at")
    )

@api_router.post("/auth/complete-registration", response_model=UserResponse)
async def complete_registration(request: Request, user: dict = Depends(get_current_user)):
    body = await request.json()
    phone_number = body.get("phone_number")
    position = body.get("position")
    province = body.get("province")
    district = body.get("district")
    facility = body.get("facility")
    
    if not phone_number or not position or not province or not district or not facility:
        raise HTTPException(status_code=400, detail="Phone number, position, province, district, and facility are required")
    
    # Validate province
    province_names = [p["name"] for p in PROVINCES]
    if province not in province_names:
        raise HTTPException(status_code=400, detail="Invalid province")
    
    # Validate district (check hardcoded + DB)
    hardcoded_districts = DISTRICTS.get(province, [])
    db_facs = await db.facilities.find({"province": province}, {"_id": 0, "district": 1}).to_list(1000)
    db_districts = [f["district"] for f in db_facs if f.get("district")]
    all_districts = set(hardcoded_districts + db_districts)
    if district not in all_districts:
        raise HTTPException(status_code=400, detail="Invalid district for selected province")
    
    # Validate facility (check hardcoded + DB)
    hardcoded_facilities = FACILITIES_BY_DISTRICT.get(district, [])
    db_facs2 = await db.facilities.find({"district": district}, {"_id": 0, "name": 1}).to_list(1000)
    db_facility_names = [f["name"] for f in db_facs2 if f.get("name")]
    all_facilities = set(hardcoded_facilities + db_facility_names)
    if facility not in all_facilities:
        raise HTTPException(status_code=400, detail="Invalid facility for selected district")
    
    # Validate position
    if position not in POSITIONS:
        raise HTTPException(status_code=400, detail="Invalid position")
    
    await db.users.update_one(
        {"user_id": user["user_id"]},
        {"$set": {"phone_number": phone_number, "position": position, "province": province, "district": district, "facility": facility}}
    )
    
    updated_user = await db.users.find_one({"user_id": user["user_id"]}, {"_id": 0})
    
    return UserResponse(
        user_id=updated_user["user_id"],
        email=updated_user["email"],
        name=updated_user["name"],
        phone_number=updated_user.get("phone_number"),
        position=updated_user.get("position"),
        province=updated_user.get("province"),
        district=updated_user.get("district"),
        facility=updated_user.get("facility"),
        area_of_allocation=updated_user.get("area_of_allocation"),
        picture=updated_user.get("picture"),
        role=updated_user.get("role", "user"),
        created_at=updated_user.get("created_at")
    )

@api_router.post("/auth/logout")
async def logout(request: Request, response: Response):
    session_token = request.cookies.get("session_token")
    if session_token:
        await db.user_sessions.delete_one({"session_token": session_token})
    
    response.delete_cookie(key="session_token", path="/")
    return {"message": "Logged out successfully"}

# ===================== DATA ROUTES =====================

@api_router.get("/facilities")
async def get_facilities():
    return {"facilities": FACILITIES}

@api_router.get("/positions")
async def get_positions():
    return {"positions": POSITIONS}

@api_router.get("/areas")
async def get_areas():
    return {"areas": AREAS_OF_ALLOCATION}

@api_router.get("/provinces")
async def get_provinces():
    return {"provinces": PROVINCES}

@api_router.get("/districts/{province}")
async def get_districts(province: str):
    # Merge hardcoded districts with DB-stored facilities' districts
    hardcoded = DISTRICTS.get(province, [])
    db_facilities = await db.facilities.find({"province": province}, {"_id": 0, "district": 1}).to_list(1000)
    db_districts = list({f["district"] for f in db_facilities if f.get("district")})
    merged = sorted(set(hardcoded + db_districts))
    return {"province": province, "districts": merged}

@api_router.get("/facilities/{district}")
async def get_facilities_by_district(district: str):
    # Merge hardcoded facilities with DB-stored facilities for this district
    hardcoded = FACILITIES_BY_DISTRICT.get(district, [])
    db_facilities = await db.facilities.find({"district": district}, {"_id": 0, "name": 1}).to_list(1000)
    db_names = [f["name"] for f in db_facilities if f.get("name")]
    merged = sorted(set(hardcoded + db_names))
    return {"district": district, "facilities": merged}

# ===================== ATTENDANCE ROUTES =====================

@api_router.post("/attendance", response_model=AttendanceRecord)
async def create_attendance(attendance: AttendanceCreate, user: dict = Depends(get_current_user)):
    if not user.get("position") or not user.get("facility"):
        raise HTTPException(status_code=400, detail="Please complete your profile first")
    
    if attendance.action not in ["login", "logout"]:
        raise HTTPException(status_code=400, detail="Action must be 'login' or 'logout'")
    
    attendance_id = f"att_{uuid.uuid4().hex[:12]}"
    
    # Use area from request, fallback to user's stored area
    area = attendance.area_of_allocation or user.get("area_of_allocation")
    
    record = {
        "attendance_id": attendance_id,
        "user_id": user["user_id"],
        "user_name": user["name"],
        "position": user["position"],
        "facility": user["facility"],
        "area_of_allocation": area,
        "action": attendance.action,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "latitude": attendance.latitude,
        "longitude": attendance.longitude,
        "synced": True,
        "offline_id": attendance.offline_id
    }
    
    await db.attendance.insert_one(record)
    
    return AttendanceRecord(**{k: v for k, v in record.items() if k != "offline_id"})

@api_router.post("/attendance/sync")
async def sync_offline_attendance(sync_data: OfflineAttendanceSync, user: dict = Depends(get_current_user)):
    """Sync offline attendance records"""
    if not user.get("position") or not user.get("facility"):
        raise HTTPException(status_code=400, detail="Please complete your profile first")
    
    synced_records = []
    
    for record in sync_data.records:
        # Check if already synced (by offline_id)
        if record.offline_id:
            existing = await db.attendance.find_one({"offline_id": record.offline_id}, {"_id": 0})
            if existing:
                synced_records.append(existing)
                continue
        
        attendance_id = f"att_{uuid.uuid4().hex[:12]}"
        
        # Use area from request, fallback to user's stored area
        area = record.area_of_allocation or user.get("area_of_allocation")
        
        new_record = {
            "attendance_id": attendance_id,
            "user_id": user["user_id"],
            "user_name": user["name"],
            "position": user["position"],
            "facility": user["facility"],
            "area_of_allocation": area,
            "action": record.action,
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "latitude": record.latitude,
            "longitude": record.longitude,
            "synced": True,
            "offline_id": record.offline_id
        }
        
        await db.attendance.insert_one(new_record)
        synced_records.append(new_record)
    
    return {"synced": len(synced_records), "records": synced_records}

@api_router.get("/attendance/me")
async def get_my_attendance(
    limit: int = 50,
    skip: int = 0,
    user: dict = Depends(get_current_user)
):
    """Get current user's attendance history"""
    records = await db.attendance.find(
        {"user_id": user["user_id"]},
        {"_id": 0}
    ).sort("timestamp", -1).skip(skip).limit(limit).to_list(limit)
    
    total = await db.attendance.count_documents({"user_id": user["user_id"]})
    
    return {"records": records, "total": total}

@api_router.get("/attendance/status")
async def get_attendance_status(user: dict = Depends(get_current_user)):
    """Get user's current attendance status (logged in or not)"""
    today = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
    
    # Get today's records
    today_records = await db.attendance.find(
        {
            "user_id": user["user_id"],
            "timestamp": {"$gte": today.isoformat()}
        },
        {"_id": 0}
    ).sort("timestamp", -1).to_list(100)
    
    if not today_records:
        return {"status": "not_logged_in", "last_action": None}
    
    last_record = today_records[0]
    
    if last_record["action"] == "login":
        return {"status": "on_duty", "last_action": last_record}
    else:
        return {"status": "off_duty", "last_action": last_record}

# ===================== ADMIN ROUTES =====================

@api_router.get("/admin/users")
async def admin_get_users(
    skip: int = 0,
    limit: int = 50,
    facility: Optional[str] = None,
    user: dict = Depends(get_admin_user)
):
    query = {}
    # Admins are scoped to their district; superusers see all
    if user.get("role") == "admin":
        admin_district = user.get("district")
        if admin_district:
            # Get all facilities in this district
            hardcoded_facs = FACILITIES_BY_DISTRICT.get(admin_district, [])
            db_facs = await db.facilities.find({"district": admin_district}, {"_id": 0, "name": 1}).to_list(500)
            db_names = [f["name"] for f in db_facs]
            all_facs = list(set(hardcoded_facs + db_names))
            if all_facs:
                query["facility"] = {"$in": all_facs}
    
    if facility:
        query["facility"] = facility
    
    users = await db.users.find(query, {"_id": 0, "password": 0}).skip(skip).limit(limit).to_list(limit)
    total = await db.users.count_documents(query)
    
    return {"users": users, "total": total}

@api_router.put("/admin/users/{user_id}")
async def admin_update_user(user_id: str, update: UserUpdate, user: dict = Depends(get_admin_user)):
    # Admins cannot edit superusers
    target = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    if target.get("role") == "superuser" and user.get("role") != "superuser":
        raise HTTPException(status_code=403, detail="Cannot edit a super user account")
    
    update_data = {k: v for k, v in update.model_dump().items() if v is not None}
    
    if not update_data:
        raise HTTPException(status_code=400, detail="No update data provided")
    
    # Admins cannot assign superuser role
    if update_data.get("role") == "superuser" and user.get("role") != "superuser":
        raise HTTPException(status_code=403, detail="Cannot assign super user role")
    
    result = await db.users.update_one({"user_id": user_id}, {"$set": update_data})
    
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="User not found")
    
    updated_user = await db.users.find_one({"user_id": user_id}, {"_id": 0, "password": 0})
    return updated_user

@api_router.put("/admin/users/{user_id}/shift")
async def admin_assign_shift(user_id: str, request: Request, user: dict = Depends(get_admin_user)):
    """Assign a shift to a user. Shift types are defined by the super user."""
    body = await request.json()
    shift_type = body.get("shift_type")  # morning, afternoon, night, four_off, custom
    custom_start = body.get("custom_start")  # for custom shifts e.g. "07:30"
    custom_end = body.get("custom_end")      # for custom shifts e.g. "16:00"
    
    valid_shifts = ["morning", "afternoon", "night", "four_off", "custom"]
    if shift_type not in valid_shifts:
        raise HTTPException(status_code=400, detail=f"Invalid shift type. Must be one of: {valid_shifts}")
    
    target = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    
    # Build shift assignment
    shift_data = {"assigned_shift": shift_type}
    if shift_type == "custom":
        if not custom_start or not custom_end:
            raise HTTPException(status_code=400, detail="Custom shift requires custom_start and custom_end times")
        shift_data["custom_shift_start"] = custom_start
        shift_data["custom_shift_end"] = custom_end
    
    await db.users.update_one({"user_id": user_id}, {"$set": shift_data})
    updated = await db.users.find_one({"user_id": user_id}, {"_id": 0, "password": 0})
    return updated

@api_router.get("/admin/shifts")
async def admin_get_shifts(user: dict = Depends(get_admin_user)):
    """Get shift configuration defined by super user (read-only for admin)"""
    config = await db.shift_config.find_one({}, {"_id": 0})
    if not config:
        config = {
            "config_id": "default",
            "morning_start": "06:00", "morning_end": "14:00",
            "afternoon_start": "14:00", "afternoon_end": "22:00",
            "night_start": "22:00", "night_end": "06:00",
            "four_off_start": "07:00", "four_off_end": "19:00",
            "grace_period_minutes": 15
        }
    return config

@api_router.get("/admin/attendance")
async def admin_get_attendance(
    skip: int = 0,
    limit: int = 100,
    facility: Optional[str] = None,
    date: Optional[str] = None,
    user_name: Optional[str] = None,
    user: dict = Depends(get_admin_user)
):
    query = {}
    
    # Admins scoped to their district
    if user.get("role") == "admin":
        admin_district = user.get("district")
        if admin_district:
            hardcoded_facs = FACILITIES_BY_DISTRICT.get(admin_district, [])
            db_facs = await db.facilities.find({"district": admin_district}, {"_id": 0, "name": 1}).to_list(500)
            db_names = [f["name"] for f in db_facs]
            all_facs = list(set(hardcoded_facs + db_names))
            if all_facs:
                query["facility"] = {"$in": all_facs}
    
    if facility:
        query["facility"] = facility
    
    if date:
        # Parse date and create range
        try:
            date_obj = datetime.fromisoformat(date.replace('Z', '+00:00'))
            start = date_obj.replace(hour=0, minute=0, second=0, microsecond=0)
            end = start + timedelta(days=1)
            query["timestamp"] = {"$gte": start.isoformat(), "$lt": end.isoformat()}
        except Exception:
            pass
    
    if user_name:
        query["user_name"] = {"$regex": user_name, "$options": "i"}
    
    records = await db.attendance.find(query, {"_id": 0}).sort("timestamp", -1).skip(skip).limit(limit).to_list(limit)
    total = await db.attendance.count_documents(query)
    
    return {"records": records, "total": total}

@api_router.get("/admin/attendance/realtime")
async def admin_get_realtime_attendance(user: dict = Depends(get_admin_user)):
    """Get real-time attendance - staff currently on duty"""
    today = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
    
    realtime_query = {"timestamp": {"$gte": today.isoformat()}}
    
    # Admins scoped to their district
    if user.get("role") == "admin":
        admin_district = user.get("district")
        if admin_district:
            hardcoded_facs = FACILITIES_BY_DISTRICT.get(admin_district, [])
            db_facs = await db.facilities.find({"district": admin_district}, {"_id": 0, "name": 1}).to_list(500)
            db_names = [f["name"] for f in db_facs]
            all_facs = list(set(hardcoded_facs + db_names))
            if all_facs:
                realtime_query["facility"] = {"$in": all_facs}
    
    today_records = await db.attendance.find(
        realtime_query,
        {"_id": 0}
    ).sort("timestamp", -1).to_list(1000)
    
    # Group by user and find those who are on duty (last action is login)
    user_status = {}
    for record in today_records:
        uid = record["user_id"]
        if uid not in user_status:
            user_status[uid] = record
    
    on_duty = [r for r in user_status.values() if r["action"] == "login"]
    
    # Get facility breakdown
    facility_counts = {}
    for record in on_duty:
        facility = record["facility"]
        facility_counts[facility] = facility_counts.get(facility, 0) + 1
    
    return {
        "on_duty_count": len(on_duty),
        "on_duty_staff": on_duty,
        "facility_breakdown": facility_counts
    }

@api_router.get("/admin/export")
async def admin_export_attendance(
    facility: Optional[str] = None,
    date: Optional[str] = None,
    format: str = "csv",
    user: dict = Depends(get_admin_user)
):
    """Export attendance records as CSV or Excel"""
    query = {}
    
    if facility:
        query["facility"] = facility
    
    if date:
        try:
            date_obj = datetime.fromisoformat(date.replace('Z', '+00:00'))
            start = date_obj.replace(hour=0, minute=0, second=0, microsecond=0)
            end = start + timedelta(days=1)
            query["timestamp"] = {"$gte": start.isoformat(), "$lt": end.isoformat()}
        except Exception:
            pass
    
    records = await db.attendance.find(query, {"_id": 0}).sort("timestamp", -1).to_list(10000)
    
    if format == "csv":
        # Generate CSV
        import csv
        from io import StringIO
        
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(["Staff Name", "Position", "Facility", "Location Type", "Action", "Timestamp", "Latitude", "Longitude"])
        
        for r in records:
            writer.writerow([
                r.get("user_name", ""),
                r.get("position", ""),
                r.get("facility", ""),
                r.get("area_of_allocation", ""),
                r.get("action", ""),
                r.get("timestamp", ""),
                r.get("latitude", ""),
                r.get("longitude", "")
            ])
        
        return Response(
            content=output.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=attendance_{date or 'all'}.csv"}
        )
    else:
        # Generate Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance"
        
        headers = ["Staff Name", "Position", "Facility", "Location Type", "Action", "Timestamp", "Latitude", "Longitude"]
        ws.append(headers)
        
        for r in records:
            ws.append([
                r.get("user_name", ""),
                r.get("position", ""),
                r.get("facility", ""),
                r.get("area_of_allocation", ""),
                r.get("action", ""),
                r.get("timestamp", ""),
                r.get("latitude", ""),
                r.get("longitude", "")
            ])
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=attendance_{date or 'all'}.xlsx"}
        )

# ===================== EMAIL BACKUP ROUTE =====================

@api_router.post("/admin/send-backup")
async def send_backup_email(
    date: Optional[str] = None,
    user: dict = Depends(get_admin_user)
):
    """Send attendance backup email"""
    if not resend.api_key:
        raise HTTPException(status_code=500, detail="Email service not configured")
    
    query = {}
    if date:
        try:
            date_obj = datetime.fromisoformat(date.replace('Z', '+00:00'))
            start = date_obj.replace(hour=0, minute=0, second=0, microsecond=0)
            end = start + timedelta(days=1)
            query["timestamp"] = {"$gte": start.isoformat(), "$lt": end.isoformat()}
        except Exception:
            pass
    
    records = await db.attendance.find(query, {"_id": 0}).sort("timestamp", -1).to_list(10000)
    
    # Build HTML table
    rows_html = ""
    for r in records:
        rows_html += f"""
        <tr>
            <td style="padding: 8px; border: 1px solid #ddd;">{r.get("user_name", "")}</td>
            <td style="padding: 8px; border: 1px solid #ddd;">{r.get("position", "")}</td>
            <td style="padding: 8px; border: 1px solid #ddd;">{r.get("facility", "")}</td>
            <td style="padding: 8px; border: 1px solid #ddd;">{r.get("area_of_allocation", "N/A")}</td>
            <td style="padding: 8px; border: 1px solid #ddd;">{r.get("action", "").upper()}</td>
            <td style="padding: 8px; border: 1px solid #ddd;">{r.get("timestamp", "")}</td>
            <td style="padding: 8px; border: 1px solid #ddd;">{r.get("latitude", "N/A")}, {r.get("longitude", "N/A")}</td>
        </tr>
        """
    
    html_content = f"""
    <html>
    <body style="font-family: Arial, sans-serif;">
        <h2 style="color: #0f766e;">V-Chron Attendance Backup</h2>
        <p>Date: {date or 'All records'}</p>
        <p>Total Records: {len(records)}</p>
        <table style="border-collapse: collapse; width: 100%;">
            <thead>
                <tr style="background-color: #0f766e; color: white;">
                    <th style="padding: 8px; border: 1px solid #ddd;">Staff Name</th>
                    <th style="padding: 8px; border: 1px solid #ddd;">Position</th>
                    <th style="padding: 8px; border: 1px solid #ddd;">Facility</th>
                    <th style="padding: 8px; border: 1px solid #ddd;">Location Type</th>
                    <th style="padding: 8px; border: 1px solid #ddd;">Action</th>
                    <th style="padding: 8px; border: 1px solid #ddd;">Timestamp</th>
                    <th style="padding: 8px; border: 1px solid #ddd;">GPS Coordinates</th>
                </tr>
            </thead>
            <tbody>
                {rows_html}
            </tbody>
        </table>
        <p style="margin-top: 20px; color: #666;">Generated by V-Chron - The Truth of Time</p>
    </body>
    </html>
    """
    
    params = {
        "from": SENDER_EMAIL,
        "to": [BACKUP_EMAIL],
        "subject": f"V-Chron Attendance Backup - {date or 'All Records'}",
        "html": html_content
    }
    
    try:
        email = await asyncio.to_thread(resend.Emails.send, params)
        return {"message": "Backup email sent successfully", "email_id": email.get("id")}
    except Exception as e:
        logger.error(f"Failed to send email: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to send email: {str(e)}")

# ===================== SUPER USER ROUTES =====================

@api_router.get("/superuser/stats")
async def superuser_stats(user: dict = Depends(get_superuser)):
    """Get dashboard statistics for super user"""
    total_users = await db.users.count_documents({})
    total_admins = await db.users.count_documents({"role": "admin"})
    total_superusers = await db.users.count_documents({"role": "superuser"})
    total_facilities_count = await db.facilities.count_documents({})
    
    today = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
    today_logins = await db.attendance.count_documents({
        "action": "login",
        "timestamp": {"$gte": today.isoformat()}
    })
    today_logouts = await db.attendance.count_documents({
        "action": "logout",
        "timestamp": {"$gte": today.isoformat()}
    })
    
    # Currently on duty
    today_records = await db.attendance.find(
        {"timestamp": {"$gte": today.isoformat()}},
        {"_id": 0}
    ).sort("timestamp", -1).to_list(5000)
    
    user_status = {}
    for record in today_records:
        uid = record["user_id"]
        if uid not in user_status:
            user_status[uid] = record
    on_duty = sum(1 for r in user_status.values() if r["action"] == "login")
    
    return {
        "total_users": total_users,
        "total_admins": total_admins,
        "total_superusers": total_superusers,
        "total_facilities": total_facilities_count if total_facilities_count > 0 else len(FACILITIES),
        "today_logins": today_logins,
        "today_logouts": today_logouts,
        "currently_on_duty": on_duty
    }

@api_router.post("/superuser/promote")
async def promote_to_superuser(request: Request, user: dict = Depends(get_superuser)):
    """Promote a user to superuser by email"""
    body = await request.json()
    email = body.get("email")
    if not email:
        raise HTTPException(status_code=400, detail="Email is required")
    
    target = await db.users.find_one({"email": email}, {"_id": 0})
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    
    await db.users.update_one({"email": email}, {"$set": {"role": "superuser"}})
    return {"message": f"User {email} promoted to superuser"}

@api_router.get("/superuser/users")
async def superuser_get_users(
    skip: int = 0,
    limit: int = 100,
    role: Optional[str] = None,
    facility: Optional[str] = None,
    search: Optional[str] = None,
    user: dict = Depends(get_superuser)
):
    """Get all users with filters (Super User only)"""
    query = {}
    if role and role != "all":
        query["role"] = role
    if facility and facility != "all":
        query["facility"] = facility
    if search:
        query["$or"] = [
            {"name": {"$regex": search, "$options": "i"}},
            {"email": {"$regex": search, "$options": "i"}}
        ]
    
    users = await db.users.find(query, {"_id": 0, "password": 0}).skip(skip).limit(limit).to_list(limit)
    total = await db.users.count_documents(query)
    return {"users": users, "total": total}

@api_router.get("/superuser/provinces")
async def superuser_get_provinces(user: dict = Depends(get_superuser)):
    """Get all provinces"""
    return {"provinces": PROVINCES}

@api_router.get("/superuser/districts")
async def superuser_get_districts(user: dict = Depends(get_superuser)):
    """Get all districts by province"""
    return {"districts": DISTRICTS}

@api_router.delete("/superuser/users/{user_id}")
async def delete_user(user_id: str, user: dict = Depends(get_superuser)):
    """Delete a user (Super User only)"""
    if user["user_id"] == user_id:
        raise HTTPException(status_code=400, detail="Cannot delete yourself")
    
    result = await db.users.delete_one({"user_id": user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="User not found")
    
    # Also delete user sessions
    await db.user_sessions.delete_many({"user_id": user_id})
    
    return {"message": "User deleted successfully"}

@api_router.post("/superuser/users/{user_id}/reset-password")
async def reset_user_password(user_id: str, data: PasswordReset, user: dict = Depends(get_superuser)):
    """Reset a user's password (Super User only)"""
    target_user = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    if not target_user:
        raise HTTPException(status_code=404, detail="User not found")
    
    hashed_password = hash_password(data.new_password)
    await db.users.update_one(
        {"user_id": user_id},
        {"$set": {"password": hashed_password}}
    )
    
    return {"message": "Password reset successfully"}

@api_router.put("/superuser/users/{user_id}/role")
async def update_user_role(user_id: str, request: Request, user: dict = Depends(get_superuser)):
    """Update user role and assigned scope (Super User only)"""
    body = await request.json()
    role = body.get("role")
    assigned_scope = body.get("assigned_scope")
    
    if role not in ["user", "admin", "superuser"]:
        raise HTTPException(status_code=400, detail="Invalid role")
    
    update_data = {"role": role}
    if assigned_scope:
        update_data["assigned_scope"] = assigned_scope
    
    result = await db.users.update_one({"user_id": user_id}, {"$set": update_data})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="User not found")
    
    updated_user = await db.users.find_one({"user_id": user_id}, {"_id": 0, "password": 0})
    return updated_user

# Facility Management
@api_router.get("/superuser/facilities")
async def get_all_facilities_admin(user: dict = Depends(get_superuser)):
    """Get all facilities with full details"""
    facilities = await db.facilities.find({}, {"_id": 0}).to_list(1000)
    if not facilities:
        # Initialize from FACILITIES_BY_DISTRICT if empty
        for district, facility_list in FACILITIES_BY_DISTRICT.items():
            province = "Central Province"  # Default for Mkushi
            for facility_name in facility_list:
                await db.facilities.insert_one({
                    "facility_id": f"fac_{uuid.uuid4().hex[:12]}",
                    "name": facility_name,
                    "district": district,
                    "province": province,
                    "created_at": datetime.now(timezone.utc).isoformat()
                })
        facilities = await db.facilities.find({}, {"_id": 0}).to_list(1000)
    return {"facilities": facilities}

@api_router.post("/superuser/facilities")
async def create_facility(facility: FacilityCreate, user: dict = Depends(get_superuser)):
    """Create a new facility"""
    existing = await db.facilities.find_one({"name": facility.name, "district": facility.district}, {"_id": 0})
    if existing:
        raise HTTPException(status_code=400, detail="Facility already exists")
    
    facility_doc = {
        "facility_id": f"fac_{uuid.uuid4().hex[:12]}",
        "name": facility.name,
        "district": facility.district,
        "province": facility.province,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    
    await db.facilities.insert_one(facility_doc)
    created = await db.facilities.find_one({"facility_id": facility_doc["facility_id"]}, {"_id": 0})
    return created

@api_router.put("/superuser/facilities/{facility_id}")
async def update_facility(facility_id: str, facility: FacilityUpdate, user: dict = Depends(get_superuser)):
    """Update a facility"""
    update_data = {k: v for k, v in facility.model_dump().items() if v is not None}
    if not update_data:
        raise HTTPException(status_code=400, detail="No update data provided")
    
    result = await db.facilities.update_one({"facility_id": facility_id}, {"$set": update_data})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Facility not found")
    
    updated = await db.facilities.find_one({"facility_id": facility_id}, {"_id": 0})
    return updated

@api_router.delete("/superuser/facilities/{facility_id}")
async def delete_facility(facility_id: str, user: dict = Depends(get_superuser)):
    """Delete a facility"""
    result = await db.facilities.delete_one({"facility_id": facility_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Facility not found")
    return {"message": "Facility deleted successfully"}

# Shift Configuration
@api_router.get("/superuser/shifts")
async def get_shift_config(user: dict = Depends(get_superuser)):
    """Get shift configuration"""
    config = await db.shift_config.find_one({}, {"_id": 0})
    if not config:
        config = {
            "config_id": "default",
            "morning_start": "06:00",
            "morning_end": "14:00",
            "afternoon_start": "14:00",
            "afternoon_end": "22:00",
            "night_start": "22:00",
            "night_end": "06:00",
            "four_off_start": "07:00",
            "four_off_end": "19:00",
            "grace_period_minutes": 15
        }
        await db.shift_config.insert_one(config)
        config = await db.shift_config.find_one({"config_id": "default"}, {"_id": 0})
    return config

@api_router.put("/superuser/shifts")
async def update_shift_config(config: ShiftConfig, user: dict = Depends(get_superuser)):
    """Update shift configuration"""
    config_data = config.model_dump()
    config_data["config_id"] = "default"
    config_data["updated_at"] = datetime.now(timezone.utc).isoformat()
    
    await db.shift_config.update_one(
        {"config_id": "default"},
        {"$set": config_data},
        upsert=True
    )
    
    return config_data

# Attendance Reports with Late/Early highlighting
@api_router.get("/superuser/attendance-report")
async def get_attendance_report(
    date: Optional[str] = None,
    province: Optional[str] = None,
    district: Optional[str] = None,
    facility: Optional[str] = None,
    user: dict = Depends(get_superuser)
):
    """Get attendance report with late/early status"""
    query = {}
    
    if facility:
        query["facility"] = facility
    elif district:
        # Get all facilities in district - merge DB + hardcoded
        db_facilities = await db.facilities.find({"district": district}, {"name": 1, "_id": 0}).to_list(500)
        db_names = [f["name"] for f in db_facilities]
        hardcoded_names = FACILITIES_BY_DISTRICT.get(district, [])
        all_fac_names = list(set(db_names + hardcoded_names))
        if all_fac_names:
            query["facility"] = {"$in": all_fac_names}
    elif province:
        # Get all facilities in province - merge DB + hardcoded districts/facilities
        db_facilities = await db.facilities.find({"province": province}, {"name": 1, "_id": 0}).to_list(2000)
        db_names = [f["name"] for f in db_facilities]
        province_districts = DISTRICTS.get(province, [])
        hardcoded_names = []
        for d in province_districts:
            hardcoded_names.extend(FACILITIES_BY_DISTRICT.get(d, []))
        all_fac_names = list(set(db_names + hardcoded_names))
        if all_fac_names:
            query["facility"] = {"$in": all_fac_names}
    
    if date:
        try:
            date_obj = datetime.fromisoformat(date.replace('Z', '+00:00'))
            start = date_obj.replace(hour=0, minute=0, second=0, microsecond=0)
            end = start + timedelta(days=1)
            query["timestamp"] = {"$gte": start.isoformat(), "$lt": end.isoformat()}
        except Exception:
            pass
    
    # Get shift config
    shift_config = await db.shift_config.find_one({}, {"_id": 0})
    if not shift_config:
        shift_config = {
            "morning_start": "06:00",
            "afternoon_start": "14:00",
            "night_start": "22:00",
            "four_off_start": "07:00",
            "grace_period_minutes": 15
        }
    
    grace_minutes = shift_config.get("grace_period_minutes", 15)
    
    records = await db.attendance.find(query, {"_id": 0}).sort("timestamp", -1).to_list(10000)
    
    # Build a map of user_id -> assigned_shift for accurate late/early calculation
    user_ids = list({r["user_id"] for r in records})
    users_data = await db.users.find(
        {"user_id": {"$in": user_ids}}, 
        {"_id": 0, "user_id": 1, "assigned_shift": 1, "custom_shift_start": 1, "custom_shift_end": 1}
    ).to_list(len(user_ids))
    user_shift_map = {u["user_id"]: u for u in users_data}
    
    # Process records to add late/early status
    processed_records = []
    for record in records:
        if record.get("action") == "login":
            timestamp = record.get("timestamp", "")
            try:
                dt = datetime.fromisoformat(timestamp.replace('Z', '+00:00'))
                time_str = dt.strftime("%H:%M")
                
                # Determine expected shift start based on user's assigned shift
                user_info = user_shift_map.get(record["user_id"], {})
                user_shift = user_info.get("assigned_shift", "morning")
                
                shift_start_map = {
                    "morning": shift_config.get("morning_start", "06:00"),
                    "afternoon": shift_config.get("afternoon_start", "14:00"),
                    "night": shift_config.get("night_start", "22:00"),
                    "four_off": shift_config.get("four_off_start", "07:00"),
                    "custom": user_info.get("custom_shift_start", "06:00")
                }
                expected_start = shift_start_map.get(user_shift, shift_config.get("morning_start", "06:00"))
                
                expected_time = datetime.strptime(expected_start, "%H:%M")
                grace_time = expected_time + timedelta(minutes=grace_minutes)
                actual_time = datetime.strptime(time_str, "%H:%M")
                
                record["assigned_shift"] = user_shift
                
                if actual_time <= expected_time:
                    record["status"] = "early"
                elif actual_time <= grace_time:
                    record["status"] = "on_time"
                else:
                    record["status"] = "late"
                    total_late_mins = int((actual_time - grace_time).total_seconds() / 60)
                    record["minutes_late"] = total_late_mins
                    hours = total_late_mins // 60
                    mins = total_late_mins % 60
                    if hours > 0:
                        record["late_display"] = f"{hours}h {mins}m"
                    else:
                        record["late_display"] = f"{mins}m"
            except Exception:
                record["status"] = "unknown"
        else:
            record["status"] = "logout"
        
        processed_records.append(record)
    
    # Calculate summary
    late_count = sum(1 for r in processed_records if r.get("status") == "late")
    early_count = sum(1 for r in processed_records if r.get("status") == "early")
    on_time_count = sum(1 for r in processed_records if r.get("status") == "on_time")
    
    return {
        "records": processed_records,
        "summary": {
            "total": len([r for r in processed_records if r.get("action") == "login"]),
            "late": late_count,
            "early": early_count,
            "on_time": on_time_count
        }
    }

# Export by scope
@api_router.get("/superuser/export")
async def superuser_export(
    province: Optional[str] = None,
    district: Optional[str] = None,
    facility: Optional[str] = None,
    date: Optional[str] = None,
    format: str = "xlsx",
    user: dict = Depends(get_superuser)
):
    """Export attendance data filtered by province/district/facility"""
    query = {}
    
    if facility:
        query["facility"] = facility
    elif district:
        db_facilities = await db.facilities.find({"district": district}, {"name": 1, "_id": 0}).to_list(500)
        db_names = [f["name"] for f in db_facilities]
        hardcoded_names = FACILITIES_BY_DISTRICT.get(district, [])
        all_fac_names = list(set(db_names + hardcoded_names))
        if all_fac_names:
            query["facility"] = {"$in": all_fac_names}
    elif province:
        db_facilities = await db.facilities.find({"province": province}, {"name": 1, "_id": 0}).to_list(2000)
        db_names = [f["name"] for f in db_facilities]
        province_districts = DISTRICTS.get(province, [])
        hardcoded_names = []
        for d in province_districts:
            hardcoded_names.extend(FACILITIES_BY_DISTRICT.get(d, []))
        all_fac_names = list(set(db_names + hardcoded_names))
        if all_fac_names:
            query["facility"] = {"$in": all_fac_names}
    
    if date:
        try:
            date_obj = datetime.fromisoformat(date.replace('Z', '+00:00'))
            start = date_obj.replace(hour=0, minute=0, second=0, microsecond=0)
            end = start + timedelta(days=1)
            query["timestamp"] = {"$gte": start.isoformat(), "$lt": end.isoformat()}
        except Exception:
            pass
    
    records = await db.attendance.find(query, {"_id": 0}).sort("timestamp", -1).to_list(10000)
    
    # Get shift config for late/early calculation
    shift_config = await db.shift_config.find_one({}, {"_id": 0}) or {"morning_start": "06:00", "grace_period_minutes": 15}
    grace_minutes = shift_config.get("grace_period_minutes", 15)
    morning_start = shift_config.get("morning_start", "06:00")
    
    if format == "csv":
        import csv
        from io import StringIO
        
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(["Staff Name", "Position", "Facility", "Location Type", "Action", "Timestamp", "Status", "Time Late", "Latitude", "Longitude"])
        
        for r in records:
            status = ""
            late_display = ""
            if r.get("action") == "login":
                try:
                    dt = datetime.fromisoformat(r.get("timestamp", "").replace('Z', '+00:00'))
                    time_str = dt.strftime("%H:%M")
                    expected_time = datetime.strptime(morning_start, "%H:%M")
                    grace_time = expected_time + timedelta(minutes=grace_minutes)
                    actual_time = datetime.strptime(time_str, "%H:%M")
                    
                    if actual_time <= expected_time:
                        status = "EARLY"
                    elif actual_time <= grace_time:
                        status = "ON TIME"
                    else:
                        status = "LATE"
                        total_late_mins = int((actual_time - grace_time).total_seconds() / 60)
                        hours = total_late_mins // 60
                        mins = total_late_mins % 60
                        late_display = f"{hours}h {mins}m" if hours > 0 else f"{mins}m"
                except Exception:
                    status = ""
            
            writer.writerow([
                r.get("user_name", ""),
                r.get("position", ""),
                r.get("facility", ""),
                r.get("area_of_allocation", ""),
                r.get("action", ""),
                r.get("timestamp", ""),
                status,
                late_display,
                r.get("latitude", ""),
                r.get("longitude", "")
            ])
        
        scope_name = facility or district or province or "all"
        return Response(
            content=output.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=attendance_{scope_name}_{date or 'all'}.csv"}
        )
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Report"
        
        headers = ["Staff Name", "Position", "Facility", "Location Type", "Action", "Timestamp", "Status", "Time Late", "Latitude", "Longitude"]
        for col_idx, header_val in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header_val)
        
        # Style header
        from openpyxl.styles import PatternFill, Font
        from openpyxl.utils import get_column_letter
        header_fill = PatternFill(start_color="FF0F766E", end_color="FF0F766E", fill_type="solid")
        header_font = Font(color="FFFFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        
        late_fill = PatternFill(start_color="FFFEE2E2", end_color="FFFEE2E2", fill_type="solid")
        early_fill = PatternFill(start_color="FFDCFCE7", end_color="FFDCFCE7", fill_type="solid")
        
        for r in records:
            status = ""
            row_fill = None
            if r.get("action") == "login":
                try:
                    dt = datetime.fromisoformat(r.get("timestamp", "").replace('Z', '+00:00'))
                    time_str = dt.strftime("%H:%M")
                    expected_time = datetime.strptime(morning_start, "%H:%M")
                    grace_time = expected_time + timedelta(minutes=grace_minutes)
                    actual_time = datetime.strptime(time_str, "%H:%M")
                    
                    if actual_time <= expected_time:
                        status = "EARLY"
                        row_fill = early_fill
                    elif actual_time <= grace_time:
                        status = "ON TIME"
                        row_fill = early_fill
                    else:
                        status = "LATE"
                        row_fill = late_fill
                        total_late_mins = int((actual_time - grace_time).total_seconds() / 60)
                        hours = total_late_mins // 60
                        mins = total_late_mins % 60
                        late_display = f"{hours}h {mins}m" if hours > 0 else f"{mins}m"
                except Exception:
                    status = ""
                    late_display = ""
            
            row = [
                r.get("user_name", "") or "",
                r.get("position", "") or "",
                r.get("facility", "") or "",
                r.get("area_of_allocation", "") or "",
                r.get("action", "") or "",
                r.get("timestamp", "") or "",
                status,
                late_display if status == "LATE" else "",
                str(r.get("latitude") or ""),
                str(r.get("longitude") or "")
            ]
            row_num = ws.max_row + 1
            for col_idx, val in enumerate(row, 1):
                clean_val = "" if val is None else val
                cell = ws.cell(row=row_num, column=col_idx, value=clean_val)
                if row_fill:
                    cell.fill = row_fill
        
        # Set column widths for readability
        col_widths = [25, 20, 30, 15, 10, 30, 10, 12, 12, 12]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        scope_name = facility or district or province or "all"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="attendance_{scope_name}_{date or "all"}.xlsx"'}
        )

# ===================== BOOTSTRAP SUPERUSER =====================

@api_router.post("/superuser/bootstrap")
async def bootstrap_superuser(request: Request):
    """Bootstrap: Promote a user to superuser. Only works if no superusers exist."""
    existing_su = await db.users.find_one({"role": "superuser"}, {"_id": 0})
    if existing_su:
        raise HTTPException(status_code=403, detail="A superuser already exists. Use the promote endpoint instead.")
    
    body = await request.json()
    email = body.get("email")
    if not email:
        raise HTTPException(status_code=400, detail="Email is required")
    
    target = await db.users.find_one({"email": email}, {"_id": 0})
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    
    await db.users.update_one({"email": email}, {"$set": {"role": "superuser"}})
    return {"message": f"User {email} has been promoted to superuser"}

# ===================== HEALTH CHECK =====================

@api_router.get("/")
async def root():
    return {"message": "V-Chron API is running", "version": "1.0.0"}

@api_router.get("/health")
async def health_check():
    return {"status": "healthy", "timestamp": datetime.now(timezone.utc).isoformat()}

# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
